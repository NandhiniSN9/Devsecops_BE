"""Generate Unit Test Case Excel report for DevSecOps Backend."""

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Test case data
TEST_CASES = [
    # main.py
    ("UTC_001", "main.py", "app_initialization", "FastAPI()", "Verify FastAPI application initialization with correct title and lifespan", 'FastAPI app with title "Overview Dashboard API"', "PASS", "Application lifespan validates settings at startup"),
    ("UTC_002", "main.py", "exception_handler", "InvalidParameterError", "Verify InvalidParameterError handler returns 400 JSON response", "JSONResponse with status_code=400, status='failed', message", "PASS", "Custom 400 error handler"),
    ("UTC_003", "main.py", "exception_handler", "AuthenticationError", "Verify AuthenticationError handler returns 401 JSON response", "JSONResponse with status_code=401, status='failed', message", "PASS", "Custom 401 error handler"),
    ("UTC_004", "main.py", "exception_handler", "NotFoundError", "Verify NotFoundError handler returns 404 JSON response", "JSONResponse with status_code=404, status='failed', message", "PASS", "Custom 404 error handler"),
    ("UTC_005", "main.py", "exception_handler", "Exception", "Verify global exception handler catches unhandled errors and logs to DB", "JSONResponse with status_code=500, status='error', async error log task created", "PASS", "Non-blocking error logging via asyncio.create_task"),
    # settings.py
    ("UTC_006", "settings.py", "validate_settings_at_startup", "Valid .env", "Verify settings validation passes with valid DATABASE_URL, TOKEN_PRIVATE_KEY, JIRA_BASE_URL, JIRA_API_TOKEN", "Settings instance returned successfully", "PASS", "All required env vars validated"),
    ("UTC_007", "settings.py", "validate_settings_at_startup", "Invalid DATABASE_URL", "Verify settings validation fails with invalid PostgreSQL URL format", "sys.exit(1) called with error message", "PASS", "Regex pattern validation for DB URL"),
    ("UTC_008", "settings.py", "validate_settings_at_startup", "Missing JIRA_BASE_URL", "Verify settings validation fails when JIRA_BASE_URL is empty", "sys.exit(1) called with error message", "PASS", "Required field validation"),
    # utils/helpers.py
    ("UTC_009", "utils/helpers.py", "normalize_project_name", '"zeb-touchpoint-pj"', "Verify normalization removes zeb- prefix and replaces hyphens with spaces", '"touchpoint pj"', "PASS", "Core normalization logic"),
    ("UTC_010", "utils/helpers.py", "normalize_project_name", '"Touchpoint PJ"', "Verify normalization lowercases and trims input", '"touchpoint pj"', "PASS", "Case-insensitive comparison"),
    ("UTC_011", "utils/helpers.py", "normalize_project_name", '"ZEB-Multi-Word-Project"', "Verify normalization handles uppercase ZEB prefix", '"multi word project"', "PASS", "Case-insensitive prefix removal"),
    ("UTC_012", "utils/helpers.py", "normalize_project_name", '""', "Verify normalization returns empty string for empty input", '""', "PASS", "Edge case: empty string"),
    ("UTC_013", "utils/helpers.py", "normalize_project_name", "None", "Verify normalization returns empty string for None input", '""', "PASS", "Edge case: None handling"),
    ("UTC_014", "utils/helpers.py", "normalize_project_name", '"my-zeb-project"', "Verify zeb- only removed from start, not middle", '"my zeb project"', "PASS", "Prefix-only removal"),
    # utils/exceptions
    ("UTC_015", "utils/exceptions/exceptions.py", "InvalidParameterError", '"bad param"', "Verify InvalidParameterError stores message correctly", "Exception with message='bad param'", "PASS", "Custom exception class"),
    ("UTC_016", "utils/exceptions/exceptions.py", "AuthenticationError", "No args", "Verify AuthenticationError uses default message", 'Exception with message containing "Authentication failed"', "PASS", "Default message handling"),
    ("UTC_017", "utils/exceptions/exceptions.py", "NotFoundError", '"Project not found"', "Verify NotFoundError stores custom message", "Exception with message='Project not found'", "PASS", "Custom exception class"),
    ("UTC_018", "utils/exceptions/error_responses.py", "build_invalid_parameter_error_response", '"Field X invalid"', "Verify 400 error response builder", "Dict with status_code=400, status='failed', message, data=[]", "PASS", "Standardized error response"),
    ("UTC_019", "utils/exceptions/error_responses.py", "build_authentication_error_response", "trace_id='abc-123'", "Verify 401 error response with trace_id prefix", "Dict with [trace_id:abc-123] in message", "PASS", "Trace ID correlation"),
    ("UTC_020", "utils/exceptions/error_responses.py", "build_unexpected_error_response", "No args", "Verify 500 error response builder", "Dict with status_code=500, status='error'", "PASS", "Generic server error response"),
    # services/overview_service.py
    ("UTC_021", "services/overview_service.py", "_validate_period", "None", "Verify None period defaults to 'last_week'", "Returns 'last_week'", "PASS", "Default period behavior"),
    ("UTC_022", "services/overview_service.py", "_validate_period", '"last_month"', "Verify valid period 'last_month' accepted", "Returns 'last_month'", "PASS", "Valid enum value"),
    ("UTC_023", "services/overview_service.py", "_validate_period", '""', "Verify empty string period raises InvalidParameterError", "InvalidParameterError with accepted values list", "PASS", "Empty string validation"),
    ("UTC_024", "services/overview_service.py", "_validate_period", '"invalid"', "Verify invalid period raises InvalidParameterError", "InvalidParameterError with accepted values list", "PASS", "Invalid enum value"),
    ("UTC_025", "services/overview_service.py", "_validate_period", '"Last_Week"', "Verify period validation is case-sensitive", "InvalidParameterError raised", "PASS", "Case-sensitive validation"),
    ("UTC_026", "services/overview_service.py", "_parse_specialization", "None", "Verify None specialization returns None (no filter)", "Returns None", "PASS", "No filter applied"),
    ("UTC_027", "services/overview_service.py", "_parse_specialization", "Valid UUID CSV", "Verify valid UUIDs are parsed correctly", "List of UUID objects", "PASS", "CSV parsing"),
    ("UTC_028", "services/overview_service.py", "_parse_specialization", '"not-a-uuid"', "Verify all invalid UUIDs returns None", "Returns None", "PASS", "Invalid UUID handling"),
    ("UTC_029", "services/overview_service.py", "_parse_specialization", "55 UUIDs", "Verify specialization IDs limited to 50", "List with 50 UUIDs", "PASS", "MAX_SPECIALIZATION_FILTER_COUNT"),
    ("UTC_030", "services/overview_service.py", "_build_tile", "current=5, comparison=0", "Verify KPI tile with increase trend", "KpiTileResponse(count=5, trend='increase', change=5)", "PASS", "Trend calculation"),
    ("UTC_031", "services/overview_service.py", "_build_tile", "current=2, comparison=5", "Verify KPI tile with decrease trend", "KpiTileResponse(count=2, trend='decrease', change=3)", "PASS", "Absolute change value"),
    ("UTC_032", "services/overview_service.py", "_build_tile", "current=5, comparison=5", "Verify KPI tile with flat trend", "KpiTileResponse(count=5, trend='flat', change=0)", "PASS", "No change detection"),
    ("UTC_033", "services/overview_service.py", "get_overview", "period=None, no KPI records", "Verify empty KPI returns zeros with null trends", "OverviewDataResponse with all counts=0, trends=None", "PASS", "Empty data handling"),
    ("UTC_034", "services/overview_service.py", "get_overview", "period='invalid'", "Verify invalid period raises error before querying repos", "InvalidParameterError raised, repos not called", "PASS", "Fail-fast validation"),
    ("UTC_035", "services/overview_service.py", "get_overview", "Multiple KPI records", "Verify KPI records aggregated across specializations", "Summed counts in OverviewMetricsResponse", "PASS", "Multi-specialization aggregation"),
    ("UTC_036", "services/overview_service.py", "_compute_status_distribution", "total=10, active=7, completed=3", "Verify percentage calculation (count/total)*100", "Active=70.0%, Completed=30.0%", "PASS", "Percentage rounding to 1 decimal"),
    ("UTC_037", "services/overview_service.py", "_compute_status_distribution", "total=0", "Verify zero total returns 0.0% for all", "All percentages = 0.0", "PASS", "Division by zero prevention"),
    # services/filter_service.py
    ("UTC_038", "services/filter_service.py", "get_filters", "Active data exists", "Verify filters returns sorted specializations, clients, statuses", "FiltersDataResponse with sorted arrays", "PASS", "Alphabetical sorting"),
    ("UTC_039", "services/filter_service.py", "get_filters", "No active data", "Verify filters returns empty arrays when no data", "FiltersDataResponse with empty arrays", "PASS", "Empty result handling"),
    ("UTC_040", "services/filter_service.py", "get_filters", "Mixed data", "Verify return type is FiltersDataResponse", "isinstance check passes", "PASS", "Type validation"),
    # services/settings_service.py
    ("UTC_041", "services/settings_service.py", "get_settings", "Valid specialization UUID", "Verify settings returned with recipients", "SettingsDataResponse with email_recipients list", "PASS", "Full settings retrieval"),
    ("UTC_042", "services/settings_service.py", "get_settings", '"not-a-uuid"', "Verify invalid UUID raises InvalidParameterError", "InvalidParameterError with 'specializationId' in message", "PASS", "UUID format validation"),
    ("UTC_043", "services/settings_service.py", "get_settings", "Non-existent UUID", "Verify missing specialization raises NotFoundError", "NotFoundError with 'Specialization not found'", "PASS", "404 handling"),
    ("UTC_044", "services/settings_service.py", "update_settings", "at_risk_threshold=10", "Verify threshold update applied", "update_setting_fields called with threshold=10", "PASS", "Partial update"),
    ("UTC_045", "services/settings_service.py", "update_settings", 'email_digest="daily"', "Verify valid schedule value accepted", "update_setting_fields called with email_digest='daily'", "PASS", "Valid schedule value"),
    ("UTC_046", "services/settings_service.py", "update_settings", 'email_digest="every_hour"', "Verify invalid schedule raises InvalidParameterError", "InvalidParameterError with 'email_digest' in message", "PASS", "Schedule validation"),
    ("UTC_047", "services/settings_service.py", "update_settings", 'action="add", email="new@co.com"', "Verify add recipient creates new record", "add_email_recipient called", "PASS", "Recipient addition"),
    ("UTC_048", "services/settings_service.py", "update_settings", "Add duplicate recipient", "Verify duplicate recipient raises InvalidParameterError", "InvalidParameterError with 'already exists'", "PASS", "Duplicate prevention"),
    ("UTC_049", "services/settings_service.py", "update_settings", 'action="remove", valid ID', "Verify remove recipient soft-deletes", "soft_delete_recipient called", "PASS", "Soft delete"),
    ("UTC_050", "services/settings_service.py", "update_settings", "Remove non-existent ID", "Verify removing missing recipient raises error", "InvalidParameterError with 'not found'", "PASS", "Missing recipient handling"),
    ("UTC_051", "services/settings_service.py", "update_settings", "Valid request", "Verify transaction committed after all updates", "commit() called once", "PASS", "Transaction management"),
    # services/servicenow_service.py
    ("UTC_052", "services/servicenow_service.py", "sync_projects", "New project (sn_project_id not in DB)", "Verify new project created with Inactive status", "create_project called, status_id=Inactive UUID", "PASS", "Project creation"),
    ("UTC_053", "services/servicenow_service.py", "sync_projects", "Existing project (sn_project_id in DB)", "Verify existing project updated, not duplicated", "update_project called, create_project not called", "PASS", "Deduplication by sn_project_id"),
    ("UTC_054", "services/servicenow_service.py", "sync_projects", "Multiple projects (mix new/existing)", "Verify mixed batch returns correct counts", "result: created=1, updated=1, total=2", "PASS", "Batch processing"),
    ("UTC_055", "services/servicenow_service.py", "sync_projects", "New project", "Verify created_by set to authenticated service account", "project.created_by = 'servicenow@zeb.co'", "PASS", "Audit trail"),
    ("UTC_056", "services/servicenow_service.py", "sync_devsecops_tickets", "Ticket with matching sn_project_id", "Verify ticket created and linked to project by ID", "create_ticket called with correct project_id", "PASS", "Step 1: ID match"),
    ("UTC_057", "services/servicenow_service.py", "sync_devsecops_tickets", "Ticket with matching project_name", "Verify project resolved by exact name (case-insensitive)", "get_project_by_name called", "PASS", "Step 2: Name match"),
    ("UTC_058", "services/servicenow_service.py", "sync_devsecops_tickets", 'Ticket name "zeb-touchpoint-pj"', "Verify project resolved by normalized name match", "get_project_by_normalized_name called, matched 'Touchpoint PJ'", "PASS", "Step 3: Normalized name match"),
    ("UTC_059", "services/servicenow_service.py", "sync_devsecops_tickets", "Ticket with matching client", "Verify project resolved by client fallback", "get_project_by_client called", "PASS", "Step 4: Client match"),
    ("UTC_060", "services/servicenow_service.py", "sync_devsecops_tickets", "Ticket with no matches", "Verify default project used as last resort", "get_default_project called", "PASS", "Step 5: Default project"),
    ("UTC_061", "services/servicenow_service.py", "sync_devsecops_tickets", "No project found (including default)", "Verify ticket skipped and counted as failed", "result: created=0, failed=1", "PASS", "Graceful skip"),
    ("UTC_062", "services/servicenow_service.py", "sync_devsecops_tickets", "Unknown specialization_name", "Verify ticket skipped when specialization not found", "result: failed=1, create_ticket not called", "PASS", "Specialization validation"),
    ("UTC_063", "services/servicenow_service.py", "sync_devsecops_tickets", "Ticket with repositories", "Verify repositories created with lead_approvers as CSV", "repository.lead_approvers = 'lead1@co,lead2@co'", "PASS", "Repository creation"),
    ("UTC_064", "services/servicenow_service.py", "sync_devsecops_tickets", "Project with is_devsecops_onboarded=False", "Verify project flagged as onboarded after ticket link", "project.is_devsecops_onboarded = True", "PASS", "Onboarding flag set"),
    ("UTC_065", "services/servicenow_service.py", "sync_devsecops_tickets", "Project already onboarded", "Verify modified_at not changed if already onboarded", "project.modified_at unchanged", "PASS", "Idempotent flag"),
    ("UTC_066", "services/servicenow_service.py", "sync_devsecops_tickets", "Existing ticket (sn_project_id match)", "Verify existing ticket updated instead of duplicated", "update_ticket called, create_ticket not called", "PASS", "Ticket upsert"),
    ("UTC_067", "services/servicenow_service.py", "sync_devsecops_tickets", "2 tickets: 1 valid, 1 invalid", "Verify partial success (valid processed, invalid skipped)", "result: created=1, failed=1, total=2", "PASS", "Partial success pattern"),
    ("UTC_068", "services/servicenow_service.py", "sync_devsecops_tickets", "Exception during processing", "Verify exception counted as failed, not crash", "result: failed=1, no unhandled exception", "PASS", "Error resilience"),
    # routes/overview_route.py
    ("UTC_069", "routes/overview_route.py", "get_overview", "GET /api/v1/overview", "Verify endpoint returns 200 with BaseResponse structure", "status_code=200, status='success', data contains metrics", "PASS", "Success response"),
    ("UTC_070", "routes/overview_route.py", "get_overview", "GET /api/v1/overview?period=last_month", "Verify period query param passed to service", "service.get_overview called with 'last_month'", "PASS", "Query param forwarding"),
    ("UTC_071", "routes/overview_route.py", "get_overview", "GET /api/v1/overview (no params)", "Verify None passed for both period and specialization", "service.get_overview called with (None, None)", "PASS", "Default params"),
    # routes/filter_route.py
    ("UTC_072", "routes/filter_route.py", "get_filters", "GET /api/v1/filters", "Verify endpoint returns 200 with Cache-Control header", "status='success', Cache-Control: max-age=3600", "PASS", "Caching header set"),
    # routes/servicenow_route.py
    ("UTC_073", "routes/servicenow_route.py", "sync_projects", "POST valid payload", "Verify endpoint returns 200 on successful sync", "status='success', message='Sync completed successfully'", "PASS", "Happy path"),
    ("UTC_074", "routes/servicenow_route.py", "sync_projects", "POST empty projects list", "Verify 422 returned for empty projects array", "422 validation error", "PASS", "min_length=1 validation"),
    ("UTC_075", "routes/servicenow_route.py", "sync_projects", "POST missing required fields", "Verify 422 returned for missing project_name", "422 validation error with field details", "PASS", "Pydantic field validation"),
    ("UTC_076", "routes/servicenow_route.py", "sync_devsecops_tickets", "POST valid payload with repos", "Verify endpoint returns 200 on successful ticket sync", "status='success', message='Sync completed successfully'", "PASS", "Happy path with repositories"),
    ("UTC_077", "routes/servicenow_route.py", "sync_devsecops_tickets", "POST with devSec_project_id alias", "Verify alias field accepted and parsed", "devsec_project_id populated in request DTO", "PASS", "Field alias support"),
    ("UTC_078", "routes/servicenow_route.py", "sync_devsecops_tickets", "POST missing specialization_name", "Verify 422 returned for missing required field", "422 validation error", "PASS", "Required field validation"),
    ("UTC_079", "routes/servicenow_route.py", "sync_devsecops_tickets", "Service raises RuntimeError", "Verify 500 returned for unhandled service error", "status='error', status_code=500", "PASS", "Error propagation"),
    # routes/settings_route.py
    ("UTC_080", "routes/settings_route.py", "get_settings", "GET /api/v1/settings/{uuid}", "Verify endpoint returns settings with recipients", "status='success', data contains setting_id", "PASS", "Settings retrieval"),
    ("UTC_081", "routes/settings_route.py", "update_settings", "PUT /api/v1/settings/manage", "Verify endpoint updates and returns new settings", "status='success', message='Settings updated successfully'", "PASS", "Settings update"),
    # routes/report_route.py
    ("UTC_082", "routes/report_route.py", "generate_reports", "POST /api/v1/reports/generate", "Verify endpoint triggers report generation", "status='success', message='Email sent successfully'", "PASS", "Cron trigger endpoint"),
    ("UTC_083", "routes/report_route.py", "generate_reports", "Service raises exception", "Verify 500 returned when report generation fails", "status='error', status_code=500", "PASS", "Error handling"),
    ("UTC_084", "routes/report_route.py", "generate_reports", "GET method", "Verify GET method not allowed", "405 Method Not Allowed", "PASS", "HTTP method validation"),
    # routes/default_route.py
    ("UTC_085", "routes/default_route.py", "health_check", "GET /health", "Verify health endpoint returns healthy status", '{"status": "healthy"}', "PASS", "Liveness probe"),
    ("UTC_086", "routes/default_route.py", "readiness_check", "GET /ready", "Verify readiness endpoint checks DB connectivity", '{"status": "ready"} with SELECT 1 query', "PASS", "Readiness probe with timeout"),
    # dtos/request validation
    ("UTC_087", "dtos/request/servicenow_request.py", "SyncProjectItemRequest", "Whitespace in sn_project_id", "Verify whitespace stripped from string fields", "sn_project_id trimmed", "PASS", "field_validator strip_whitespace"),
    ("UTC_088", "dtos/request/servicenow_request.py", "SyncProjectRequest", "Empty projects list", "Verify min_length=1 constraint on projects", "ValidationError raised", "PASS", "List minimum length"),
    ("UTC_089", "dtos/request/settings_request.py", "SettingsUpdateRequest", "at_risk_threshold=0", "Verify ge=1 constraint rejects 0", "ValidationError raised", "PASS", "Minimum value constraint"),
    ("UTC_090", "dtos/request/settings_request.py", "EmailRecipientAction", "Invalid email format", "Verify EmailStr rejects invalid email", "ValidationError raised", "PASS", "Email format validation"),
    # dtos/response validation
    ("UTC_091", "dtos/response/base_response.py", "BaseResponse", "message > 256 chars", "Verify max_length=256 constraint on message", "ValidationError raised", "PASS", "Message length limit"),
    ("UTC_092", "dtos/response/base_response.py", "BaseResponse", 'status="unknown"', "Verify Literal constraint rejects invalid status", "ValidationError raised", "PASS", "Enum validation"),
    ("UTC_093", "dtos/response/overview_response.py", "KpiTileResponse", "count=-1", "Verify ge=0 constraint rejects negative count", "ValidationError raised", "PASS", "Non-negative constraint"),
    ("UTC_094", "dtos/response/overview_response.py", "KpiTileResponse", 'trend="invalid"', "Verify Literal constraint rejects invalid trend", "ValidationError raised", "PASS", "Trend enum validation"),
    ("UTC_095", "dtos/response/overview_response.py", "StatusBreakdownItemResponse", "percentage=100.1", "Verify le=100.0 constraint", "ValidationError raised", "PASS", "Percentage upper bound"),
    # services/dependencies.py
    ("UTC_096", "services/dependencies.py", "get_db_session", "Database session", "Verify async session yielded and closed", "AsyncSession instance", "PASS", "DI for database"),
    ("UTC_097", "services/dependencies.py", "get_overview_service", "Service factory", "Verify OverviewService created with injected repo", "OverviewService instance", "PASS", "DI chain: session → repo → service"),
    ("UTC_098", "services/dependencies.py", "get_servicenow_service", "Service factory", "Verify ServiceNowService created with injected repo", "ServiceNowService instance", "PASS", "DI for ServiceNow"),
    # PeriodEnum
    ("UTC_099", "dtos/request/overview_request.py", "PeriodEnum", '"last_week"', "Verify valid enum value accepted", "PeriodEnum.LAST_WEEK", "PASS", "String enum"),
    ("UTC_100", "dtos/request/overview_request.py", "PeriodEnum", '"LAST_WEEK"', "Verify case-sensitive rejection", "ValueError raised", "PASS", "Case sensitivity"),
]

PREPARED_BY = "Nandhini"
EXECUTED_DATE = date(2026, 5, 18).strftime("%m/%d/%y")
EXECUTED_BY = "Nandhini"


def generate_excel():
    """Generate the unit test case Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    # Headers
    headers = [
        "Unit Test Case ID", "Module Name", "Sub Module Name",
        "Input Request", "Unit Test Case Description", "Expected Result",
        "Prepared By", "Executed Date - Before Deployment",
        "Executed By - Before Deployment", "Result", "Status", "Comments"
    ]

    # Style
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Write data
    for row_idx, tc in enumerate(TEST_CASES, 2):
        tc_id, module, sub_module, input_req, description, expected, status, comments = tc
        row_data = [
            tc_id, module, sub_module, input_req, description, expected,
            PREPARED_BY, EXECUTED_DATE, EXECUTED_BY, expected, status, comments
        ]
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    # Column widths
    col_widths = [16, 32, 28, 30, 55, 50, 12, 18, 18, 50, 8, 40]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    output_path = "DevSecOps_BE_Unit_Test_Cases.xlsx"
    wb.save(output_path)
    print(f"Generated: {output_path}")
    print(f"Total test cases: {len(TEST_CASES)}")


if __name__ == "__main__":
    generate_excel()
